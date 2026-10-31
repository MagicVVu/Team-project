"""Build the standalone Garin work record using the bundled document runtime."""
import json,pathlib
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=pathlib.Path(__file__).resolve().parents[1]
def load(n):return json.loads((ROOT/'reports'/n).read_text(encoding='utf-8'))
s=load('summary.json');cats=load('分功能结果.json');u=load('统计不确定性.json');ext=load('extremes.json')['records']
d=Document();sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 st=d.styles[name];st.font.name='Microsoft YaHei';st.font.color.rgb=RGBColor(0,0,0);st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'微软雅黑')
st=d.styles['Normal'];st.font.size=Pt(10.5);st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1.15
d.styles['Title'].font.size=Pt(23);d.styles['Title'].paragraph_format.space_after=Pt(10)
d.styles['Heading 1'].font.size=Pt(16);d.styles['Heading 1'].paragraph_format.space_before=Pt(8);d.styles['Heading 1'].paragraph_format.space_after=Pt(8)
d.styles['Heading 2'].font.size=Pt(12);d.styles['Heading 2'].paragraph_format.space_before=Pt(8)
for border in list(d.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
d.styles['Subtitle'].font.italic=False
d.core_properties.title='P8 6 DeepSeek基线测试工作记录';d.core_properties.author='Garin-cyber';d.core_properties.subject='真实模型基线与失败证据交接';d.core_properties.comments='人工复核未签字。'
def p(text,style=None):return d.add_paragraph(text,style)
def h(text):d.add_heading(text,1)
def table(headers,rows,widths):
 t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for i,w in enumerate(widths):t.columns[i].width=Cm(w)
 for i,x in enumerate(headers):t.rows[0].cells[i].text=x
 for row in rows:
  for c,x in zip(t.add_row().cells,row):c.text=str(x)
 for j,row in enumerate(t.rows):
  for i,c in enumerate(row.cells):
   c.width=Cm(widths[i]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcPr=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:w'),'75');el.set(qn('w:type'),'dxa');mar.append(el)
   tcPr.append(mar);b=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');b.append(el)
   tcPr.append(b);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'183B56' if j==0 else 'F1F5F8' if j%2 else 'FFFFFF');tcPr.append(shade)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(1);para.paragraph_format.line_spacing=1.05
    if i>0:para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in para.runs:r.font.size=Pt(9.5);r.font.bold=j==0;r.font.color.rgb=RGBColor(255,255,255) if j==0 else RGBColor(0,0,0)
  trPr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 d.add_paragraph().paragraph_format.space_after=Pt(1)
 return t

d.add_paragraph('P8 6 DeepSeek基线测试工作记录','Title')
p('Garin-cyber  |  2026年10月31日  |  交接给AI及项目成员','Subtitle')
p('本轮在隔离副本中测试现有AI代码与DeepSeek，Jev未启用。已执行719题次，654通过、65失败；最后1题次在发出请求前被5元预算保护拦截。以下成绩属于本地AI流程的严格自动判定，人工标注复核仍待完成。')
h('本轮主要数据')
table(['指标','本轮结果'],[
 ['目标源码','733d339  自主循环  single-object-1'],['输入规模','120条单轮加20个六轮场景  共240个输入'],['执行与自动通过率','719题次  654除以719等于90.96%'],['实际API请求','正式766次  含小批量共773次'],['每题平均API次数','1.065次  包含合法零调用'],['每题平均耗时','1.279秒'],['每题P50及P95','1.224秒  2.394秒'],['每题最长耗时','4.177秒'],['单次API平均耗时','1.041秒'],['每题平均费用','0.00103989元'],['单次API平均费用','0.00097609元'],['正式批次用量估算','0.74768412元  非平台账单'],['本轮保守费用','4.957662元  含小批量  上限5元']], [7,10])
p('每题耗时包含本地处理、实际模型请求和记录开销，不含App网络与业务数据库。单次API耗时从实际HTTP发出计到完整响应读完；本次非流式调用，不报告首字耗时。')
p('费用按真实Token、缓存和执行时段价格估算。预算采用峰值未缓存价格保守记账，最后一次最坏预留会越界，因此停止；没有清零旧账本或绕过额度。')

d.add_page_break();h('失败集中在哪些功能')
p('共65次自动失败，覆盖14类功能。下表按题次计数，同一输入重复失败不应登记成多个独立缺陷；完整预期、实际回复和证据位置已逐条保留。')
table(['功能','执行题次','失败题次'],[(c['category'],c['executed'],c['FAIL']) for c in cats if c['FAIL']], [10,3.5,3.5])
d.add_heading('优先交给AI负责人的问题',2)
p('第一，用户明确说“还没选哪一件”时仍生成移动草稿，以及过期指代继续查询，需优先核对对象选择和引用有效性。草稿未自动写入，不能把这点当作选择正确。')
p('第二，唯一对象的整件消耗和整件丢弃共有18次失败，出现继续追问数量、比例或拒绝的情况。请对照模型计划的amount_mode、字段落地和逐件规则定位，当前不直接归责单一组件。')
p('第三，部分多轮查询已经返回事实，却因重复计划而停在stopped；部分错误码与预期不同。这些严格FAIL继续保留，需分别讨论终止状态、错误码契约与事实正确性。')
p('规划资料缺失使用故意不可用的合成接口，验证降级状态，不代表正式服务器一定未实现。未发现已确认的样本标注错误；有争议的状态与错误码交由团队复核，未擅自改成PASS。')

d.add_page_break();h('调用次数与极端值')
table(['每题实际API次数','题次'],[['0次',105],['1次',463],['2次',150],['3次',1]], [10,7])
p('已证实当前链路可以零次、一次、两次或三次调用。两次不一定是在判断上次回答是否正确：136题属于核对后再单独finish，11题追加读取后再决策，3题携带校验反馈再决策。单凭次数不能说明每一步作用相同。')
d.add_heading('最慢且调用最多的题',2)
p('r1-p86-q22-2：“把牛奶开封并移到厨房”。整题4.177秒、3次API、费用0.002577元，自动PASS表示正确触发边界保护，并不表示已执行两个动作。网络请求累计3.697秒，占整题88.5%；日志出现UNGROUNDED_FIELD和MULTIPLE_WRITES，后续请求携带前轮错误反馈。')
d.add_heading('最贵的题',2)
p('r1-p86-m09-2-t6：“重新核对这件牛奶现在的数量和位置”。整题2.375秒、2次API、费用0.00306964元，自动FAIL，停止原因为repeated_plan。两次输入合计7981 Token，缓存命中70.6%；长上下文和重复决策增加了费用。')
d.add_heading('如何理解短路径与波动',2)
small=next(e for e in ext if e['kind']=='最小值')
p(f"最短路径为{small['execution_id']}，耗时{small['seconds']:.6f}秒，实际API为{small['api_calls']}次。这反映本地规则路径，不能用来说明DeepSeek推理只需这一时间。Top5全部保留，并用Q3加1.5倍IQR标注高值；高值排名不一定都是统计异常。")
p('239个输入完成三遍，其中199个三遍均通过，176个业务结果一致。最后一个输入只有两遍，均值按两遍计算。完整六轮场景执行59次，47次六轮全部通过，另1次场景因预算未完成。')
p('按34个表达或场景族整体重采样的描述性95%区间为84.81%至95.52%。重复运行不算新增独立样本，合成集结果不能当成真实用户总体正确率。客户端日志也不能确认供应商排队、推理和公网延迟各占多少。')

d.add_page_break();h('交付文件与接续工作')
p('全部本轮脚本、冻结用例、原始证据和报告存于 tests/8.6基线测试/P8.6_DeepSeek基线_20261031_733d339。给队友先阅读的副本存于 tests/交接文件夹/2026-10-31_P8.6基线测试。')
d.add_heading('建议阅读顺序',2)
for text in ['先打开Excel总览，再看每题均值。每题有独立编号，包含遍数、最短和最长耗时、费用、调用次数。','遇到问题查逐题记录、API明细与极端值页；JSONL保留全部实际结果，原始runs目录保留请求、响应、上下文、trace及状态。','失败交接清单逐项列出预期、实际回复、失败断言与证据文件行号；分类是定位线索，未确认的根因继续待复核。']:
 p(text,'List Number')
d.add_heading('本轮已完成的保护与检查',2)
p('在源码副本与虚拟家庭库存中执行，每遍重建相同状态；未连接业务服务或真实用户数据库。2379份受保护原文件及314份副本哈希无变化。计量和源码测试最终55项通过，初始工具错误与修复记录保留。所有已执行题的库存和采购事实未被写入。')
p('密钥保存在仓库外，公开附件不含鉴权头。实际请求只允许访问DeepSeek。源码、依赖、提示词、配置和冻结数据哈希可复核；测试结果没有通过改预期“修好”。')
d.add_heading('下一步由团队接续',2)
for text in ['AI负责人先处理对象选择、引用有效性、整件操作，再统一循环结束状态与缺数据错误码。Garin随后按同编号复测，保留原失败。','人工复核用例标注及回复可读性；未签字内容继续保持pending。本轮不登记为App、后端或真实数据库联合验收通过。','接入Jev后冻结同样数据和评分，分别统计DeepSeek、Jev及合计费用；尽量同期重跑DeepSeek对照，记录模型别名、网络、缓存和价格变化。','若用本轮数据调优，另备未见评估集。新的付费运行另立批次和预算，不能覆盖本轮证据。']:
 p(text,'List Number')
p('价格依据：DeepSeek官方定价页，2026年10月31日复核。https://api-docs.deepseek.com/zh-cn/quick_start/pricing/。本轮实际返回模型别名记录于API明细；别名不等于不可变的模型版本。')
p('人工复核状态：待团队确认。此记录不代替Garin本人或其他成员签字。')
out=ROOT/'reports/Garin_P8.6基线测试工作记录.docx';d.save(out);print(out)
