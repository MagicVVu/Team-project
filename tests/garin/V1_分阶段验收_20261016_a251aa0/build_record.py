"""Create the separate Chinese Word handoff; no human sign-off is implied."""
import json,pathlib,sys
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
root=pathlib.Path(sys.argv[1]).resolve();out=pathlib.Path(sys.argv[2]).resolve()
android=json.loads((root/'results/Android结果.json').read_text(encoding='utf8'))
doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.7);sec.bottom_margin=Cm(1.7);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
for style in doc.styles:
    if style.type==1:
        style.font.name='Microsoft YaHei';style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei');style.font.color.rgb=RGBColor(0,0,0)
        style.paragraph_format.space_after=Pt(5);style.paragraph_format.line_spacing=Pt(15)
        pp=style._element.get_or_add_pPr();grid=OxmlElement('w:snapToGrid');grid.set(qn('w:val'),'0');pp.append(grid)
for b in list(doc.styles.element.iter(qn('w:pBdr'))):b.getparent().remove(b)
doc.styles['Normal'].font.size=Pt(10)
doc.styles['Title'].font.size=Pt(20);doc.styles['Title'].paragraph_format.space_after=Pt(9)
for name,size in [('Heading 1',14),('Heading 2',11.5)]:
    doc.styles[name].font.size=Pt(size);doc.styles[name].font.bold=True;doc.styles[name].paragraph_format.space_before=Pt(9);doc.styles[name].paragraph_format.space_after=Pt(5)
doc.core_properties.title='Garin V1 分阶段测试工作记录';doc.core_properties.subject='2026年10月16日第一批离线复测与联调待办';doc.core_properties.author='garin-cyber 测试工作记录'
def p(t,style=None):return doc.add_paragraph(t,style)
def h(t):return doc.add_heading(t,1)
def sub(t):return doc.add_heading(t,2)
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    pr=t._tbl.tblPr;b=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        x=OxmlElement('w:'+side);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');b.append(x)
    pr.append(b)
    for i,values in enumerate([headers]+rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        for j,(cell,text) in enumerate(zip(cells,values)):
            cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;cell.text=str(text)
            for par in cell.paragraphs:
                par.paragraph_format.space_before=Pt(3);par.paragraph_format.space_after=Pt(3);par.paragraph_format.line_spacing=Pt(13)
                if j==0 or (len(headers)>2 and j==1):par.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for run in par.runs:run.font.size=Pt(9.3);run.bold=(i==0)
            if i==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7EEF5');cell._tc.get_or_add_tcPr().append(sh)
        trpr=t.rows[i]._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
        if i==0:repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
    p('')
    return t
def newpage(title):doc.add_page_break();h(title)
p('Garin V1 分阶段测试工作记录','Title')
p('第一批离线复测与联调待办  2026年10月16日')
p('工作归属：garin-cyber。面向AI、前后端与UI成员交接；自动化执行和证据已整理，人工样本复核仍以原文件签字状态为准。')
p('本轮已完成P0至P5现阶段可离线执行的检查、后端临时SQLite测试、Android调试构建及17项单元测试。Lint仍受离线依赖阻塞，App与真实PostgreSQL闭环尚未执行，因此不能宣布完整V1验收通过。')
h('基线和实际成绩')
p('固定基线：dev分支 a251aa033586e915cb49e35bbd242efccb358de6。全部运行于隔离源码副本，AI与云端后端各用独立Python环境，默认rules。')
table(['测试组','实际结果','说明'],[
['AI源码首轮','462通过 2失败','副本缺文件及P5测试误报，均保留'],
['AI源码完整复跑','464通过','补齐同版本部署模板；未改断言'],
['P1 P2 P3 P5规则集','各90加60通过','各阶段分别统计，不当真人样本相加'],
['Garin专项','22项通过','对应24条观察证据'],
['云端后端本地','40项通过','临时SQLite；另有9条弃用警告'],
['库存采购参考流程','1项通过','总库存3盒、幂等与冲突检查'],
['Android调试构建','通过','未安装APK，未连接服务'],
['Android单元测试','17项通过','JUnit证据已独立保存'],
['Android Lint','BLOCKED','离线缓存缺少Compose依赖'],
],[3.6,4,9.4])
p('各组复用了部分相同功能和数据，不能合成一个总通过率。规则评测中的model_calls代表提取次数，本轮付费模型调用为0。')
p('未访问18080正式服务，未请求18081或连接真实用户数据库。原684个范围外跟踪文件、78个旧Garin文件及410个副本输入的哈希均保持一致。')
newpage('各阶段的检查和旧问题复测')
table(['阶段','本轮检查重点'],[
['P0 基础操作','原150条、扩展96条及32条旧流程复用；另查中文数量、完整品牌、注入与多操作。草稿生成前后库存保持一致。'],
['P1 生命周期','开封、整批移动、部分及全部丢弃；唯一批次、多批次选择和零库存终止批次分别准备。检查预览、澄清与版本，不把预览当正式提交。'],
['P2 风险','期限类型、缺保存条件、到期风险、事实和规则版本；冰箱位置不当温度，绿色结果不保证安全。'],
['P3 预测补货','固定时钟及完整观察日、消费去重、非消费事件排除、待购扣减和缺数据降级。'],
['P4 历史问题','原84次程序问题逐项对照，保留所有旧断言与FAIL。117项P4源码用例包含在464项内。'],
['P5 确认与任务','识别99盒改2盒；口误、否定与转述先澄清；未确认输入不提取；不支持的公开续接在HTTP前拒绝。'],
],[3.1,13.9])
sub('两个可直接复现的参考结果')
p('P3：固定2026年10月15日，28个完整观察日消费7盒、库存2盒，得到日均0.25盒、剩余8天、参考日期10月23日。按14日覆盖和2日提前量计算，已有待购1盒时，新建议1盒。')
p('后端：新增2盒牛奶，消耗0.5盒，计划采购2盒，实买1.5盒。实际旧批1.5盒、新批1.5盒，总计3盒。重复消耗不再扣减，旧版本操作返回409，重复采购不再入库，其他家庭访问返回404。此项只验证临时SQLite中的接口行为。')
sub('旧评测器实际结果')
table(['入口','PASS','FAIL','REVIEW','BLOCKED'],[
['P0','260','0','18','0'],['P1','209','19','18','32'],['旧Android快照网关','210','18','18','32']
],[5,3,3,3,3])
p('原84次程序问题中，81条满足原断言；另外3条仍为FAIL：P1的p0-41-2返回BATCH_TERMINAL，旧Android网关的ext-013和ext-015受未注册采购品限制。它们属于现有契约边界，包含在上述37条FAIL内，不重复相加。')
newpage('失败分类和证据处理')
sub('测试程序问题')
p('P5旧测试在整个JSON中搜索99，随机request_id含99时会失败。受控复现中，改动只有UUID：不含99时通过，含99时失败；两次上传文字均为添加2盒牛奶，均无原始识别字段。首轮失败与诊断原样保留，未改队友测试代码。建议AI负责人按text和字段检查重写此断言，并用含99的固定ID防止回归。')
sub('版本契约差异和未支持功能')
p('P1保留19条旧FAIL，旧Android快照网关保留18条旧FAIL。内容包括生命周期功能新增后的状态差异、候选批次表示差异，以及旧网关采购品和单位限制。完整编号、原预期、实际输出和证据行号已逐条导出；不强改正确答案，也不把安全拒绝自动算成通过。')
p('AI草稿正式确认、开封移动丢弃事务及撤销，P3规划数据接口、公开item_key续接和依据持久化仍需后端支持。P5内部来源字段不能直接加进现有公开消息体。当前隔离启动脚本默认P1，不能用它验收P2至P5。')
sub('测试环境问题')
p('隔离副本首次漏拷部署模板，补齐同源文件后464项完整复跑通过。Android缓存复制和中文路径问题仅修复本轮适配器；临时V盘别名已解除，队友构建文件未改。Lint缺少3项Compose desktop离线依赖，保留BLOCKED，待补齐缓存后单独复跑。')
sub('历史模型失败和样本标注')
p('P5历史模型失败p5-42-1继续保留，本轮未复测真实模型。后续先说明费用和模型版本，再单独建模型复测轮次；本轮规则通过不能消除历史模型FAIL。')
p('原150条辅助复核中148条核心标签可保留，1条漏标位置、1条意图歧义。人工确认以文件为准，未代签。创可贴快用完了仍应先澄清，未明确操作意图如何编码需由团队确认。各入口的18条REVIEW也保持原状态。')
sub('证据定位')
p('失败清单.md便于成员阅读；results目录提供37条旧FAIL、155条旧非通过记录、84次问题逐项对照和24条专项观察。evidence目录保存原始输出、HTTP与前后状态、trace ID、JUnit、依赖清单及失败日志。无HTTP的直接对象测试明确记为空，不补造状态码。')
p('开发方P4验证器把3条契约边界记为resolved，并硬编码“含未提交P4修改”。该字段不代表本轮改了源码；以baseline.json及source-integrity.json为准。本轮报告单独保留3条原FAIL。')
newpage('第二批联调条件和团队交接')
p('联调待办已准备18项BLOCKED和1项样本REVIEW。每项都有数据前提、操作步骤、预期、对接人和解除阻塞条件，尚未执行的项目没有登记为通过。')
table(['对接方向','需提供或确认','后续验证'],[
['inHuman 前后端','18081健康isolated_test=true；独立数据库证明、账号、家庭、部署提交号','登录退出和缓存隔离；库存采购闭环；冲突与重复请求；真实PostgreSQL并发和回滚'],
['MagicVVu AI','实际运行入口、P阶段、引擎、提交号；P3与P5接口能力','任务读取与澄清；草稿不写事实；P2风险和P3预测依据；P5编辑文字与超时恢复'],
['summerdong2006-dot UI','专用设备及动画可操作版本','拉绳短拉取消、连续开关、离场点击、切页恢复、语音失败转手输和迟到回调'],
['Garin 与团队','人工标签和预期契约确认','保持旧FAIL与REVIEW，按新部署版本另开证据轮次'],
],[3.5,6.2,7.3])
sub('联调开始前的门槛')
p('先核对隔离服务、数据库与AI版本，任一隔离证明不足就停止写操作。使用专用模拟器或测试设备，不清空已有个人数据；真实行锁、事务与并发必须用独立PostgreSQL验证。源码更新后另开轮次，不把成绩混入本基线。')
sub('交付位置与使用顺序')
p('全部成果位于：D:\\Team_Project\\Team-project\\tests\\garin\\V1_分阶段验收_20261016_a251aa0')
p('先看README.md了解范围和复现方法，再看失败清单.md及联调待办.md。精确数据在results目录，原始证据从证据索引.md查找。本Word可单独发送给其他成员记录本轮工作。')
p('本轮名称建议用“Garin V1分阶段测试 第一批离线复测”。不建议命名为产品V1-P1，因为本轮覆盖P0至P5，会与AI成员的P1生命周期版本混淆。')
sub('本轮交接结论')
p('离线可执行部分已有实测结果和可追溯证据；Lint与正式链路条件不足，仍保留待办。未进行Git提交或推送，未代签人工复核，未修改队友负责的文件。下一批按隔离环境实际部署版本继续，补齐App和数据库闭环后再作正式验收结论。')
out.parent.mkdir(parents=True,exist_ok=True);doc.save(out);print(str(out))
