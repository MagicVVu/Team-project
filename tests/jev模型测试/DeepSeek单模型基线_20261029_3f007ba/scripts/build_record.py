"""Build the standalone Chinese Word work record from actual evidence only."""
import json
from datetime import datetime,timezone
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT

ROUND=Path(__file__).resolve().parents[1]
def read(path): return json.loads((ROUND/path).read_text(encoding='utf-8'))
def pct(value): return '未测' if value is None else f'{value*100:.2f}%'
def num(value,digits=3): return '未测' if value is None else f'{value:.{digits}f}'
def label(name):
    if name=='p84-regression':return 'P8.4 旧回归'
    if name=='p85-regression':return 'P8.5 六轮回归'
    if name.startswith('new-single-r'):return '新增单轮 第'+name[-1]+'遍'
    if name.startswith('new-multi-r'):return '新增六轮 第'+name[-1]+'遍'
    return name
def main():
    summary=read('reports/summary.json'); delivery=read('reports/delivery-status.json'); issues=read('reports/issues.json')
    formal=[r for r in summary['runs'] if r['run'].startswith('formal-')]
    smoke=[r for r in summary['runs'] if r['run'].startswith('smoke-')]
    doc=Document(); sec=doc.sections[0]
    sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
    for style in doc.styles:
        if style.type==1:
            style.font.name='Microsoft YaHei';style.font.color.rgb=RGBColor(0,0,0)
            style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
            style.paragraph_format.space_after=Pt(5);style.paragraph_format.line_spacing=1.15
    for border in list(doc.styles.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
    doc.styles['Normal'].font.size=Pt(10)
    doc.styles['Title'].font.size=Pt(21)
    for name,size in [('Heading 1',14),('Heading 2',11.5)]:
        doc.styles[name].font.size=Pt(size);doc.styles[name].font.bold=True
        doc.styles[name].paragraph_format.space_before=Pt(9);doc.styles[name].paragraph_format.space_after=Pt(5)
    doc.core_properties.author='garin-cyber 测试记录'
    doc.core_properties.title='DeepSeek 单模型基线测试工作记录'
    doc.core_properties.created=datetime.now(timezone.utc);doc.core_properties.modified=datetime.now(timezone.utc)
    def p(text): return doc.add_paragraph(text)
    def h(text): return doc.add_heading(text,1)
    def page(title): doc.add_page_break();h(title)
    def table(headers,data,widths):
        t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
        for column,width in zip(t.columns,widths):column.width=Cm(width)
        borders=OxmlElement('w:tblBorders')
        for side in ['top','left','bottom','right','insideH','insideV']:
            b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D3D3D3');borders.append(b)
        t._tbl.tblPr.append(borders)
        for i,values in enumerate([headers]+data):
            cells=t.rows[0].cells if i==0 else t.add_row().cells
            for j,(cell,text) in enumerate(zip(cells,values)):
                cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;cell.text=str(text)
                for par in cell.paragraphs:
                    par.paragraph_format.space_before=Pt(3);par.paragraph_format.space_after=Pt(3);par.paragraph_format.line_spacing=1.05
                    for run in par.runs: run.font.size=Pt(9);run.bold=i==0
                if i==0:
                    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EAEAEA');cell._tc.get_or_add_tcPr().append(sh)
            trpr=t.rows[i]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
            if i==0:trpr.append(OxmlElement('w:tblHeader'))
        p('')
    doc.add_paragraph('DeepSeek 单模型基线测试工作记录','Title')
    p('2026年10月29日　工作归属 garin-cyber　供 AI 与前后端成员交接')
    p(f"本轮正式计划1200轮，实际已有{delivery['recorded_formal_turns']}轮记录，另有{delivery['not_executed_formal_turns']}轮未执行。Jev未启用；P8.4当前默认链路与P8.5独立实现分别统计。真实模型批次启用DeepSeek，明确控制路径可以零调用；离线流程检查另列。")
    p('本轮发现真实语义失败，也发现旧评分遗漏的计算问题和部分新标注前提问题。因此不能只用旧评测的高通过率宣布全部功能正确。')
    p('源码固定为3f007ba。使用隔离副本、合成家庭和内存库存，未连接业务服务器或真实用户数据库，也未修改队友源码。人工样本复核仍为pending，没有代签。')
    p('业务测试时钟固定为10月15日10时，日期预期按该时钟计算；真实调用和计费使用10月29日实际时间。历史测试包未因本轮新基线自动全部完成真实模型复测。')
    h('正式批次执行情况')
    batchrows=[]
    for run in formal:
        for name,b in run['batches'].items():
            c=b['counts'];batchrows.append([label(name),f"{c.get('PASS',0)} / {c.get('FAIL',0)} / {c.get('REVIEW',0)}",b['actual_http_attempts'],pct(b['turn_accuracy'])])
    table(['批次','通过 失败 待复核','HTTP次数','自动评分'],batchrows,[5.6,4.8,2.3,4.3])
    p('自动评分的分母为已执行PASS和FAIL，REVIEW另列。旧回归保留原评分，其通过率不能视为所有字段和计算均正确；第二页列出独立补充核对。新增用例重复三次仍只有240轮独立输入。')
    budget=delivery['budget']
    p(f"含小批量检查的实际HTTP请求共{budget['all_live_http']}次。本轮新增保守费用{budget.get('new_conservative_cost_cny','未知')}元，上限5元，费用保护核对{'通过' if budget.get('cap_respected') else '需检查'}。供应商账单尚未核对。")
    if delivery['not_executed_formal_turns']:p('未执行项目详见not-executed.json，中止原因见正式批次complete.json。未执行不能登记为通过。')
    page('正确性和评分证据')
    p('每条新用例保存原文、虚拟库存、上下文、原预期、模型原始响应、解析结果、最终结果和逐项断言。新增写入草稿核对动作、对象、数量、单位、版本、日期和位置；查询核对资料和计算依据。')
    table(['检查角度','实际口径'],[
        ['单轮与多轮','单轮逐项满足预期才通过；完整多轮场景要求全部轮次通过。'],
        ['模型与保护机制','原始理解输出与最终草稿分别保存；模型提议错误而规则拦截时保留两层证据。'],
        ['澄清','应问未问和多余询问分别统计；多余询问是自动代理指标，仍需人工确认。'],
        ['失败与争议','超时、格式错误等已执行异常保留；有争议的标注列REVIEW，不自动改成PASS。'],
        ['置信区间','按场景组进行10000次bootstrap，固定随机种子；重复表达不当作独立真人样本。'],
    ],[3.2,13.8])
    auditrows=[]
    for run in formal:
        for name,b in run['batches'].items():
            a=b['legacy_additional_audit']
            if a['covered']:auditrows.append([label(name),a['covered'],a['pass'],a['original_pass_but_extra_failed']])
    table(['旧评测补充审核','覆盖条数','补充通过','旧PASS但补充失败'],auditrows,[6.2,2.6,2.8,5.4])
    p('补充审核不覆盖所有旧场景的完整语义；原始成绩不覆盖写回。不能把“生成了草稿”或“没有修改库存”单独作为正确的充分条件。')
    stability=summary['stability']
    p(f"完成三次重复记录的用例{stability['cases_with_three_repeats']}条；三次均通过{stability['all_three_pass']}条；业务语义输出一致{stability['same_semantic_output']}条。比较时去除随机ID和自然连接语，逐条状态及耗时保留。")
    p('本轮数据均为合成样本。旧保留集已被开发消费，按回归集解释。新增集在首次模型调用前冻结，未根据模型失败调整提示词；沿用项目类别，不能据此认定对真实用户的泛化效果。')
    page('耗时调用和费用')
    timerows=[]
    for run in formal:
        for name,b in run['batches'].items():
            t=b['ai_turn_latency_ms'];price=num(float(b['known_calculated_cost_cny'])) if b['known_calculated_cost_cny'] is not None else '未知'
            if b['cost_unknown_or_inexact_requests']:price+=' 加待核算'
            timerows.append([label(name),num(t['mean']/1000 if t['mean'] is not None else None),num(t['p95']/1000 if t['p95'] is not None else None),b['actual_http_attempts'],b['zero_wire_turns'],price])
    table(['批次','均值 秒','P95 秒','HTTP','零调用轮','用量估算 元'],timerows,[5.4,2.1,2.1,1.8,2.4,3.2])
    p('耗时包括理解、只读计算和表达，不包含手机网络、业务数据库保存或页面渲染。记录工具的少量开销未扣除；单请求与各阶段耗时另见summary.json。非流式请求不报告首字耗时。')
    p('实际HTTP请求数包含失败请求，与程序阶段计数区分。超时、限流、空回复、格式错误及备用表达均单独统计；没有自动付费重试。')
    p('价格按2026年10月29日官方页面复核。峰时每百万Token：缓存命中输入0.04元、未命中输入2元、输出8元；空闲时段为一半。本地预算始终按峰时未命中保守结算，未知用量保持预留。报告按缓存和时段给出用量估算；缺字段或跨时段不伪造精确费用。')
    p('官方价格来源 https://api-docs.deepseek.com/zh-cn/quick_start/pricing/')
    p('请求和响应的模型名逐条保存。若服务只返回deepseek-flash别名，不能声称获得了不可变的内部构建版本；官方页面对应版本与实际响应字段分开记录。模型参数与原程序一致，未指定的采样参数标记为供应商默认。')
    p('小批量检查单独统计，正式重复批次也不混成一个独立样本总分。按用量成本、保守预算和供应商实扣三种口径分别保留。')
    page('问题交接与复测要求')
    groups={}
    for issue in issues:
        if issue['run'].startswith('formal-'):
            key=issue.get('category') or issue.get('status');groups[key]=groups.get(key,0)+1
    table(['问题类型或定位线索','记录数'],[['字段或计算补充审核 根因待核对' if k=='additional field or calculation check; needs root-cause review' else k,v] for k,v in groups.items()],[14,3])
    p('记录数包含不同重复批次和补充审核。同一用例可能产生多条记录，不能相加当作独立缺陷数。')
    p('新用例中的5条风险绿色预期与当前规则审核条件不符，原标注和首次结果保留，继续待复核。采购用例的可用库存和已审核风险前提也需要团队确认；单纯算出覆盖量不能替代这些前提。')
    p('规划查询暴露了批次b1与商品milk的选择映射及旧夹具兼容问题。当前证据不足以直接等同于线上缺陷，需AI负责人核对当前接口资料。测试没有改队友代码，也没有改成只要返回文字就通过。')
    p('例如p84-22-1把多操作误当成新任务；p85-08-1第5轮用户已要求登记两盒，模型仍按否定处理，第6轮状态查询受到连带影响。p84-24-2改回当前位置冰箱未生成草稿，属于标注或版本契约待确认，原FAIL保留。')
    p('逐条预期和输出见issues.json与对应cases.jsonl，定位说明见问题定位.md。自动归因尚未获团队确认，自然表达仍待人工评估。')
    h('我本轮交付的内容')
    p('已交付源码与数据冻结清单、隔离运行脚本、请求和费用记录工具、逐条结果、失败及待复核清单、重复性统计、未执行清单和本份Word工作记录。测试脚本自检及源码相关单元检查证据保存在evidence目录。')
    p(f"最终完整性核对：范围外{delivery['integrity']['outside_owned_checked']}个文件与副本{delivery['integrity']['snapshot_checked']}个输入文件，字节校验{'一致' if delivery['integrity']['source_unchanged'] else '需检查'}。")
    h('下一轮比较要求')
    p('接入Jev后固定相同主版本、用例、初始库存、测试时钟、评分和并发条件；同时报告Jev自身调用和费用，不能只统计DeepSeek减少的部分。若修复映射或调整标注，应同时重跑DeepSeek对照，不能把这些修复计作Jev的提升。单列缓存、时段、模型别名和网络变化，App和数据库联调另行验收。')
    p(f'材料位置 tests/jev模型测试/{ROUND.name}/；交接入口 tests/交接文件夹/2026-10-29_DeepSeek单模型基线/。')
    path=ROUND/'reports/Garin DeepSeek单模型基线测试工作记录.docx';doc.save(path)
    print(str(path))
if __name__=='__main__':main()
